"""论文技术定稿及团队交接包。AI辅助：Codex/OpenAI；版本、发布日期未核实。"""
from __future__ import annotations

import csv
import hashlib
import io
import json
from pathlib import Path
import re
import tarfile
import zipfile

from .documents import make_html, render_pdf
from .io import PROJECT_ROOT, read_json, sha256, verify_sources, write_json
from .paper import collect_content, verify_evidence

DELIVERY=PROJECT_ROOT/'交付材料'
FINAL_MD=PROJECT_ROOT/'docs/论文/C题论文_技术定稿.md'
PAPER_PDF=DELIVERY/'论文/C题论文_技术定稿.pdf'


def editable_disclosure(text: str, path: Path) -> None:
    from docx import Document
    from docx.shared import Pt, Mm
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc=Document()
    section=doc.sections[0]
    section.page_width=Mm(210);section.page_height=Mm(297)
    section.top_margin=Mm(20);section.bottom_margin=Mm(20)
    section.left_margin=Mm(18);section.right_margin=Mm(18)
    style=doc.styles['Normal'];style.font.name='SimSun';style.font.size=Pt(11)
    style.element.rPr.rFonts.set(qn('w:eastAsia'),'宋体')
    def clean(value):
        return re.sub(r'\[([^]]+)\]\([^)]+\)',r'\1',value).replace('**','').replace('`','')
    lines=text.splitlines();index=0
    while index<len(lines):
        line=lines[index].strip()
        if line.startswith('|'):
            rows=[]
            while index<len(lines) and lines[index].lstrip().startswith('|'):
                cells=[clean(c.strip()) for c in lines[index].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?',c) for c in cells):
                    rows.append(cells)
                index+=1
            table=doc.add_table(rows=0,cols=len(rows[0]));table.style='Table Grid'
            for cells in rows:
                row=table.add_row()
                for cell,value in zip(row.cells,cells):
                    cell.text=value
            header=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(header)
            continue
        if line.startswith('#'):
            level=len(line)-len(line.lstrip('#'))
            doc.add_heading(clean(line.lstrip('#').strip()),level=min(level,3))
        elif line:
            doc.add_paragraph(clean(line))
        index+=1
    for key in ['author','last_modified_by','comments','subject','keywords','category']:
        setattr(doc.core_properties,key,'')
    doc.core_properties.title='AI工具使用详情（工作稿）'
    path.parent.mkdir(parents=True,exist_ok=True)
    doc.save(path)


def build_documents() -> dict:
    fingerprint=verify_evidence()
    DELIVERY.mkdir(parents=True,exist_ok=True)
    PAPER_PDF.parent.mkdir(parents=True,exist_ok=True)
    tokens,ledger=collect_content()
    tokens['SERIES_CELLS']=str(read_json(PROJECT_ROOT/'data/processed/component_curves.json')['series_cells'])
    from .remediation_report import paper_fragments
    tokens.update(paper_fragments())
    from .extension_report import paper_fragments as extension_fragments
    tokens.update(extension_fragments())
    template=PROJECT_ROOT/'paper/final.template.md'
    text=template.read_text(encoding='utf-8')
    for key,value in tokens.items():
        text=text.replace('@@'+key+'@@',value)
    assert not re.search(r'@@[A-Z_0-9]+@@',text)
    FINAL_MD.write_text(text,encoding='utf-8')
    title=text.splitlines()[0].lstrip('# ')
    html_path=PAPER_PDF.with_suffix('.html')
    make_html(text,FINAL_MD,html_path,title=title,paper=True)
    paper=render_pdf(html_path,PAPER_PDF,paper=True)
    ai_source=PROJECT_ROOT/'docs/AI工具使用详情.md'
    ai_text=ai_source.read_text(encoding='utf-8')
    ai_docx=DELIVERY/'AI工具使用详情.docx'
    editable_disclosure(ai_text,ai_docx)
    ai_html=DELIVERY/'AI工具使用详情.html'
    make_html(ai_text,ai_source,ai_html,title='AI工具使用详情（工作稿）')
    ai_pdf=DELIVERY/'AI工具使用详情.pdf'
    ai=render_pdf(ai_html,ai_pdf)
    build={'status':'technical_release_for_team_handoff','competition_submission_ready':False,
           'team_independent_work':'not_started_as_confirmed_by_user','manual_revision_and_review':'not_started',
           'paper':paper,'ai_disclosure':ai,'ai_disclosure_status':'incomplete_working_document',
           'metrics':len(ledger),'new_scientific_experiments':True,
           'new_experiments_scope':'Separate stage8 and stage10 experiments; RC preferred for voltage comparison; original main evidence retained.',
           'source_hashes':{**fingerprint,**{p:sha256(PROJECT_ROOT/p) for p in ['paper/final.template.md','paper/style.css','docs/AI工具使用详情.md','src/bev/documents.py','src/bev/delivery.py','src/bev/remediation_report.py','records/remediation_summary.json']},
                            **{p.relative_to(PROJECT_ROOT).as_posix():sha256(p) for p in (PROJECT_ROOT/'records/experiments/remediation_v1').rglob('*') if p.is_file()},
                            **{p.relative_to(PROJECT_ROOT).as_posix():sha256(p) for p in (PROJECT_ROOT/'records/experiments/extension_v1').rglob('*') if p.is_file()},
                            **{name:sha256(PROJECT_ROOT/name) for name in ['src/bev/extension.py','src/bev/extension_report.py','src/bev/voltage_rc.py','configs/extension_protocol.json','records/extension_summary.json']}},
           'output_hashes':{p.relative_to(PROJECT_ROOT).as_posix():sha256(p) for p in [FINAL_MD,PAPER_PDF,html_path,ai_docx,ai_pdf,ai_html]}}
    write_json(PROJECT_ROOT/'records/final_paper_build.json',build)
    return {k:build[k] for k in ['status','competition_submission_ready','paper','ai_disclosure']}


def trajectory_archive(path: Path, *, full: bool) -> dict:
    path.parent.mkdir(parents=True,exist_ok=True)
    entries=[]
    with tarfile.open(path,'w:xz',preset=3) as archive:
        for source in sorted((PROJECT_ROOT/'results/timeseries').glob('*.csv')):
            data=source.read_bytes();original_sha=hashlib.sha256(data).hexdigest()
            selected=None
            if not full and source.name.startswith('durability_'):
                reader=csv.reader(io.StringIO(data.decode('utf-8-sig')))
                header=next(reader)
                selected=[k for k in ['time_s','dt_s','speed_m_s','soc','battery_c','cabin_c','battery_w'] if k in header]
                indexes=[header.index(k) for k in selected]
                output=io.StringIO(newline='');writer=csv.writer(output,lineterminator='\n')
                writer.writerow(selected)
                writer.writerows([row[i] for i in indexes] for row in reader)
                data=output.getvalue().encode('utf-8-sig')
            name=source.relative_to(PROJECT_ROOT).as_posix()
            item=tarfile.TarInfo(name=name);item.size=len(data);item.mtime=0
            archive.addfile(item,io.BytesIO(data))
            entries.append({'path':name,'bytes':len(data),'sha256':hashlib.sha256(data).hexdigest(),
                            'source_full_csv_sha256':original_sha,'columns_subset':selected})
    return {'mode':'full_original_csv' if full else 'full_other_csv_and_selected_columns_for_six_long_traces',
            'files':entries,'archive_sha256':sha256(path),'archive_bytes':path.stat().st_size,
            'note':'Every retained value is an original CSV string. No resampling or rounding. Full long traces can be regenerated with reproduce.py long.'}


def expected_results() -> dict:
    metrics=[]
    def add(name,source,keys):
        value=read_json(PROJECT_ROOT/source)
        for key in keys:
            value=value[key]
        metrics.append({'name':name,'source':source,'keys':keys,'value':float(value),'relative_tolerance':1e-8,'absolute_tolerance':1e-8})
    for key in ['consumption_kwh_100km','soc_final','battery_final_c']:
        add('scenario_A_'+key,'results/tables/scenario_A_summary.json',[key])
    evaluation=read_json(PROJECT_ROOT/'results/tables/calibration_evaluation.json')
    for index,row in enumerate(evaluation['group_metrics']):
        if row['variant']=='joint_calibrated':
            for key in ['energy_mape_pct','soc_rmse_macro_pp']:
                add(row['role']+'_'+key,'results/tables/calibration_evaluation.json',['group_metrics',index,key])
    control=read_json(PROJECT_ROOT/'results/tables/control_evaluation.json')
    for index,row in enumerate(control['nominal']):
        if row['strategy']=='optimized':
            add(f"control_{row['ambient_c']:g}_saving",'results/tables/control_evaluation.json',['nominal',index,'heat_corrected_saving_pct'])
    durability=read_json(PROJECT_ROOT/'results/tables/durability_evaluation.json')
    for index,row in enumerate(durability['comparisons']):
        if row['parameter_set']=='calibrated_v2_provisional' and row['repetitions']==3 and row['parking_s']==7200:
            add(f"parking_{row['ambient_c']:g}_saving",'results/tables/durability_evaluation.json',['comparisons',index,'raw_saving_pct'])
    add('nonlinear_refits_accepted','results/tables/nonlinear_reliability.json',['accepted'])
    add('scientific_tests','results/tables/verification.json',['tests_run'])
    return {'note':'Numerical tolerances, not PDF byte identity or a new independent-data validation. all reuses frozen expensive results.', 'metrics':metrics}


def build_packages() -> dict:
    verify_evidence()
    verify_sources()
    payload_dir=PROJECT_ROOT/'records/package_payload'
    compact=payload_dir/'compact.tar.xz';full=payload_dir/'full.tar.xz'
    compact_info=trajectory_archive(compact,full=False)
    full_info=trajectory_archive(full,full=True)
    sources=[]
    for folder,suffix in [('src','.py'),('checks','.py'),('scripts','.py'),('configs','.json'),('docs','.md'),('paper','')]:
        sources.extend(p for p in (PROJECT_ROOT/folder).rglob('*') if p.is_file() and (not suffix or p.suffix==suffix) and '__pycache__' not in p.parts)
    for folder in ['data/external','data/processed','results/tables','results/figures','records/experiments','records/literature','portal']:
        sources.extend(p for p in (PROJECT_ROOT/folder).rglob('*') if p.is_file() and '__pycache__' not in p.parts)
    excluded_records={'run_manifest.json','portal_build.json','development_video.json','delivery_packages.json','delivery_verification.json','delivery_runtime_verification.json'}
    sources.extend(p for p in (PROJECT_ROOT/'records').glob('*.json') if p.name not in excluded_records)
    sources.extend((PROJECT_ROOT/'records').glob('*.md'))
    sources.extend((PROJECT_ROOT/'results').glob('*.md'))
    sources.extend([PROJECT_ROOT/'records/requirements-lock.txt',PROJECT_ROOT/'run.py',PROJECT_ROOT/'reproduce.py',PROJECT_ROOT/'AGENTS.md',PROJECT_ROOT/'C题开发日志.md'])
    sources.extend(PROJECT_ROOT.glob('requirements*.txt'))
    # 不把旧浏览器绝对路径、预览图片、开发视频或已生成的压缩包带进支撑包。
    sources=sorted(set(sources))
    disclosure={DELIVERY/'AI工具使用详情.pdf':'AI工具使用详情.pdf',DELIVERY/'AI工具使用详情.docx':'AI工具使用详情.docx'}
    common={p.relative_to(PROJECT_ROOT).as_posix():p for p in sources}
    common.update({name:p for p,name in disclosure.items()})
    readme=(PROJECT_ROOT/'docs/交付与复现说明.md').read_bytes()
    expected=json.dumps(expected_results(),ensure_ascii=False,indent=2).encode('utf-8')
    result={}
    for label,archive,trajectory_info in [('support',compact,compact_info),('team',full,full_info)]:
        files=dict(common)
        if label=='team':
            files.update({r['project_path']:PROJECT_ROOT/r['project_path'] for r in verify_sources()})
            files['C题项目管理.html']=PROJECT_ROOT/'C题项目管理.html'
            files['交付材料/论文/C题论文_技术定稿.pdf']=PAPER_PDF
            files['C题开发日志.md']=PROJECT_ROOT/'C题开发日志.md'
        generated={'README.md':readme,'EXPECTED_RESULTS.json':expected,
                   'payload/trajectories.json':json.dumps(trajectory_info,ensure_ascii=False,indent=2).encode('utf-8')}
        files['payload/trajectories.tar.xz']=archive
        entries=[{'path':name,'bytes':p.stat().st_size,'sha256':sha256(p)} for name,p in sorted(files.items())]
        entries += [{'path':name,'bytes':len(data),'sha256':hashlib.sha256(data).hexdigest()} for name,data in sorted(generated.items())]
        manifest={'type':label,'status':'team_handoff_technical_release_not_submission_ready','files':sorted(entries,key=lambda r:r['path']),
                  'official_inputs_included':label=='team','ai_disclosure_complete':False,'team_independent_work':'not_started',
                  'trajectory_mode':trajectory_info['mode']}
        generated['PACKAGE_MANIFEST.json']=json.dumps(manifest,ensure_ascii=False,indent=2).encode('utf-8')
        destination=DELIVERY/('C题_支撑材料.zip' if label=='support' else 'C题_同学复现包.zip')
        prefix='C题复现工程/'
        with zipfile.ZipFile(destination,'w',compression=zipfile.ZIP_DEFLATED,compresslevel=9) as output:
            for name,p in sorted(files.items()):
                output.write(p,prefix+name,compress_type=zipfile.ZIP_STORED if name.endswith('.xz') else zipfile.ZIP_DEFLATED)
            for name,data in sorted(generated.items()):
                output.writestr(prefix+name,data)
        if label=='support' and destination.stat().st_size>20_000_000:
            raise ValueError(f'支撑包仍超过20MB：{destination.stat().st_size}')
        with zipfile.ZipFile(destination) as zipped:
            assert zipped.testzip() is None
        result[label]={'path':destination.relative_to(PROJECT_ROOT).as_posix(),'bytes':destination.stat().st_size,
                       'sha256':sha256(destination),'files':len(entries)+1,'official_inputs_included':label=='team',
                       'trajectory_mode':trajectory_info['mode'],'submission_ready':False}
    write_json(PROJECT_ROOT/'records/delivery_packages.json',result)
    (DELIVERY/'SHA256SUMS.txt').write_text('\n'.join(f"{r['sha256']}  {Path(r['path']).name}" for r in result.values())+'\n'+f"{sha256(PAPER_PDF)}  论文/{PAPER_PDF.name}\n",encoding='utf-8')
    return result
