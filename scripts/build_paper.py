#!/usr/bin/env python3
# AI-assisted: OpenAI Codex / OpenAI, 2026-10-18.
# 具体模型快照、版本及发布日期未提供；研究内容与提交合规须由队员独立复核。
"""Generate anonymous editable manuscripts and PDFs from frozen real outputs."""
from __future__ import annotations

import argparse
import csv
from datetime import datetime, timezone
import hashlib
import json
import math
import os
from pathlib import Path
import re
import subprocess
import sys
import tempfile
import xml.etree.ElementTree as ET

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "manuscript"
OUT = ROOT / "outputs/paper"
TITLE = "基于官方路线数据的街区交通数字孪生与有界源端协同治理"
PROFILES = {"primary": "主设置", "higher_demand": "需求加倍", "severe_event": "更强慢行", "longer_event": "事件延长",
            "stronger_event": "更强慢行", "severe": "更强慢行", "longer": "事件延长"}
OD_NAMES = {"north_to_south": "北→南", "south_to_north": "南→北", "west_to_east": "西→东", "east_to_west": "东→西",
            "west_to_hospital_road": "西→医院外围", "east_to_hospital_road": "东→医院外围"}


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def workspace(path):
    path = Path(path).resolve()
    if not path.is_relative_to(ROOT):
        raise ValueError("路径必须位于当前工作区")
    return path


def load(path):
    return json.loads(workspace(path).read_text(encoding="utf-8"))


def n(value, digits=2):
    if value is None:
        return "—"
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, (int, float)):
        return f"{value:,.{digits}f}".rstrip("0").rstrip(".") if digits else f"{value:,.0f}"
    return str(value)


def table(headers, rows):
    return "\n".join(["| " + " | ".join(headers) + " |", "|" + "---|" * len(headers),
                      *("| " + " | ".join(map(str, row)) + " |" for row in rows)])


def make_network_figure(catalog):
    from PIL import Image, ImageDraw, ImageFont
    image = Image.new("RGB", (1600, 1150), "white")
    draw = ImageDraw.Draw(image)
    font_path = "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"
    font = ImageFont.truetype(font_path, 42)
    small = ImageFont.truetype(font_path, 40)
    nodes = {r["id"]: r["xy_m"] for r in catalog["network"]["nodes"]}
    edges = {r["id"]: r for r in catalog["network"]["edges"]}
    xs, ys = [p[0] for p in nodes.values()], [p[1] for p in nodes.values()]
    scale = min(1320 / (max(xs) - min(xs)), 870 / (max(ys) - min(ys)))
    cx, cy = (min(xs) + max(xs)) / 2, (min(ys) + max(ys)) / 2
    def xy(p):
        return (800 + (p[0] - cx) * scale, 565 - (p[1] - cy) * scale)
    def line(ids, color, width):
        for identity in ids:
            edge = edges[identity]
            draw.line([xy(nodes[edge["from"]]), xy(nodes[edge["to"]])], fill=color, width=width)
    line(edges, "#c4cbd0", 3)
    demand = next(d for d in catalog["demand_ods"] if d["id"] == "west_to_hospital_road")
    other = next(c for c in catalog["route_candidates"] if c["demand_od_id"] == demand["id"] and c["edge_ids"] != demand["baseline_edge_ids"])
    line(other["edge_ids"], "#30699f", 8)
    line(demand["baseline_edge_ids"], "#c78029", 5)
    for ids,color in ((other["edge_ids"],"#30699f"),(demand["baseline_edge_ids"],"#c78029")):
        for index,identity in enumerate(ids):
            edge=edges[identity]
            a,b=xy(nodes[edge["from"]]),xy(nodes[edge["to"]])
            dx,dy=b[0]-a[0],b[1]-a[1]
            length=math.hypot(dx,dy)
            if index%4 or length<25:
                continue
            ux,uy=dx/length,dy/length
            x,y=(a[0]+b[0])/2,(a[1]+b[1])/2
            draw.polygon([(x+9*ux,y+9*uy),(x-7*ux+7*uy,y-7*uy-7*ux),(x-7*ux-7*uy,y-7*uy+7*ux)],fill=color)
    event = ["e000184", "e000083", "e000084", "e000085", "e000086", "e000087"]
    line(event, "#a42337", 10)
    endpoint = xy(nodes[edges[demand["baseline_edge_ids"][-1]]["to"]])
    draw.ellipse((endpoint[0]-9, endpoint[1]-9, endpoint[0]+9, endpoint[1]+9), fill="#3e3d70")
    draw.text((endpoint[0]+20, endpoint[1]-10), "医院外围接驳点", fill="#1f2b35", font=small)
    draw.text((endpoint[0]+20, endpoint[1]+38), "（入口未核验）", fill="#1f2b35", font=small)
    draw.text((55, 22), "上海瑞金片区：已观测有向子图与西侧任务候选", fill="#1f2b35", font=font)
    draw.text((55, 77), "BD09局部投影；灰线为已观测分段，非完整路网", fill="#56616b", font=small)
    for i, (color, label) in enumerate((("#c78029", "默认路线"), ("#30699f", "经思南路备选"), ("#a42337", "事件段68.735m"))):
        x = 70 + i*470
        draw.line((x, 1090, x+60, 1090), fill=color, width=8)
        draw.text((x+74, 1061), label, fill="#1f2b35", font=small)
    draw.line((1450, 180, 1450, 120), fill="#333333", width=3)
    draw.polygon([(1450, 110), (1443, 124), (1457, 124)], fill="#333333")
    draw.text((1429, 78), "北N", fill="#333333", font=small)
    draw.line((70, 1010, 70+200*scale, 1010), fill="#333333", width=3)
    draw.text((70, 1020), "200 m（模型坐标）", fill="#333333", font=small)
    target = OUT / "figures/network.png"
    target.parent.mkdir(parents=True, exist_ok=True)
    image.save(target)
    return target


def render_svg(source):
    source = workspace(source)
    root = ET.parse(source).getroot()
    def size(key, default):
        value = root.get(key, str(default))
        return round(float(re.sub(r"[^0-9.]", "", value) or default))
    width, height = size("width", 1400), size("height", 800)
    target = OUT / "figures" / (source.stem + ".png")
    target.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix=".figure-browser-", dir=OUT) as temporary:
        command = ["/usr/bin/google-chrome", "--headless", "--no-sandbox", "--disable-gpu", "--disable-background-networking", "--no-first-run",
                   f"--user-data-dir={temporary}", f"--window-size={width},{height}", "--hide-scrollbars", f"--screenshot={target}", source.as_uri()]
        result = subprocess.run(command, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, timeout=60)
        if result.returncode or not target.exists():
            raise RuntimeError("现有浏览器无法在本环境转换SVG；请用已有PNG或获准后重新执行")
    return target


def plot_canvas(width, height):
    from PIL import Image, ImageDraw, ImageFont
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    path = "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"
    return image, draw, ImageFont.truetype(path, 42), ImageFont.truetype(path, 40)


def make_validation_figure(validation):
    image, draw, title, font = plot_canvas(1500, 1120)
    rows = list(csv.DictReader((ROOT / validation["artifacts"]["predictions.csv"]["path"]).open(encoding="utf-8")))
    draw.text((50, 22), "跨日留出：18日平台ETA与17日冻结状态", font=title, fill="#172033")
    colors = {"observed": "#202d39", "shared_affine_eta": "#aa4c32", "per_od_training_median": "#3477a0"}
    od_ids = ["north_to_south", "south_to_north", "west_to_hospital_road", "east_to_hospital_road"]
    start = datetime.fromisoformat("2026-10-18T11:36:00+08:00").timestamp()
    for index, od in enumerate(od_ids):
        left, top = 95 + (index % 2)*730, 120 + (index // 2)*465
        width, height = 580, 330
        draw.text((left, top-45), OD_NAMES[od] + "（秒）", font=title, fill="#172033")
        def point(t, y):
            return (left + (t-start)/7200*width, top+height - y/800*height)
        for y in range(0, 801, 200):
            py = point(start,y)[1]
            draw.line((left,py,left+width,py),fill="#e2e6e9",width=1)
            draw.text((left-55,py-14),str(y),font=font,fill="#54606b")
        for minute,label in ((0,"11:36"),(40,"12:16"),(80,"12:56"),(120,"13:36")):
            px = point(start+minute*60,0)[0]
            draw.text((px-30,top+height+10),label,font=font,fill="#54606b")
        selected = [r for r in rows if r["phase"]=="cross_day_holdout" and r["route_pair_id"]==od]
        observed = [r for r in selected if r["method"]=="shared_affine_eta"]
        for method in ("shared_affine_eta", "per_od_training_median", "observed"):
            values = observed if method=="observed" else [r for r in selected if r["method"]==method]
            values.sort(key=lambda r:r["requested_at"])
            coords = [point(datetime.fromisoformat(r["requested_at"]).timestamp(),float(r["observed_platform_eta_s"] if method=="observed" else r["predicted_duration_s"])) for r in values]
            draw.line(coords,fill=colors[method],width=4)
            if method=="observed":
                for x,y in coords:
                    draw.ellipse((x-4,y-4,x+4,y+4),fill=colors[method])
    for index,(method,label) in enumerate((("observed","平台ETA快照"),("per_od_training_median","OD训练中位数"),("shared_affine_eta","冻结共享仿射"))):
        x = 130+index*450
        draw.line((x,1060,x+65,1060),fill=colors[method],width=5)
        draw.text((x+80,1040),label,font=font,fill="#172033")
    target = OUT / "figures/validation_holdout.png"
    image.save(target)
    return target


def make_experiment_figure():
    image,draw,title,font = plot_canvas(1500,950)
    draw.text((50,22),"高需求S3：A/B真实仿真快照与源端等待",font=title,fill="#172033")
    draw.text((50,65),"横轴为仿真秒；事件与需求均为假设，未重建现场事故",font=font,fill="#54606b")
    frames = {p:load(ROOT / f"outputs/final_experiments/runs/higher_demand_S3_{p}/trajectory.json")["frames"] for p in ("A","B")}
    queue = load(ROOT / "outputs/final_experiments/runs/higher_demand_S3_B/source_queue_trace.json")
    colors = {"A":"#2d6e9d","B":"#a94a2c"}
    for index,(label,upper) in enumerate((("在网车辆数（辆，含任务）",120),("B源端控制等待数（辆）",20))):
        left,top,width,height = 100,165+index*370,1280,250
        draw.text((left,top-40),label,font=title,fill="#172033")
        def point(t,value):
            return (left+t/1800*width,top+height-value/upper*height)
        for value in (0,upper/2,upper):
            py=point(0,value)[1]
            draw.line((left,py,left+width,py),fill="#e2e6e9",width=1)
            draw.text((left-55,py-13),n(value,0),font=font,fill="#54606b")
        for t in (0,300,600,900,1200,1500,1800):
            x=point(t,0)[0]
            draw.text((x-20,top+height+10),str(t),font=font,fill="#54606b")
        for t,event in ((600,"事件开始"),(720,"任务发布"),(900,"控制窗口结束")):
            x=point(t,0)[0]
            draw.line((x,top,x,top+height),fill="#aeb7bf",width=1)
            if index==0:
                draw.text((x+5,top+(0 if t==600 else 35 if t==720 else 70)),event,font=font,fill="#54606b")
        if index==0:
            for p in ("A","B"):
                draw.line([point(f["time_s"],len(f["vehicles"])) for f in frames[p]],fill=colors[p],width=4)
        else:
            draw.line([point(f["time_s"],f["source_waiting"]) for f in queue],fill=colors["B"],width=4)
    for i,p in enumerate(("A","B")):
        x=520+i*250
        draw.line((x,890,x+70,890),fill=colors[p],width=5)
        draw.text((x+85,870),"策略"+p,font=font,fill="#172033")
    target = OUT / "figures/scenario_timeseries.png"
    image.save(target)
    return target


def font(run, chinese="宋体", size=12, bold=False):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:eastAsia"), chinese)
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")


def setup_document():
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2.3), Cm(2.1)
    section.left_margin, section.right_margin = Cm(2.4), Cm(2.4)
    section.footer_distance = Cm(1.0)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Pt(24)
    for name in ("Heading 1", "Heading 2"):
        style = document.styles[name]
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.keep_with_next = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.first_line_indent = Pt(0)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.title = TITLE
    document.core_properties.subject = "匿名研究稿"
    document.core_properties.keywords = "交通数字孪生;GIS;路径治理;应急协同"
    return document


def paragraph(document, text, heading=0, centered=False):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
    p = document.add_paragraph(style=f"Heading {heading}" if heading else None)
    if heading or centered:
        p.paragraph_format.first_line_indent = Pt(0)
    if heading == 1 or centered:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if heading:
        p.paragraph_format.space_before = Pt(10 if heading == 1 else 6)
        p.paragraph_format.space_after = Pt(6)
    font(p.add_run(text), "黑体" if heading == 1 else "宋体", 14 if heading == 1 else 12, bool(heading))
    return p


def add_table(document, lines):
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.shared import Pt
    cells = [[v.strip() for v in row.strip().strip("|").split("|")] for row in lines]
    cells = [row for row in cells if not all(re.fullmatch(r"[-: ]+", value) for value in row)]
    tbl = document.add_table(rows=0, cols=len(cells[0]))
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, values in enumerate(cells):
        row = tbl.add_row()
        trpr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        trpr.append(cant_split)
        if i == 0:
            repeat = OxmlElement("w:tblHeader")
            trpr.append(repeat)
        for cell, value in zip(row.cells, values):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
            font(p.add_run(value), bold=(i == 0))
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_equation(document, text):
    """Use real DOCX sub/superscript runs for the source's math indices."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    pattern = re.compile(r"(_\{[^}]*\}|_\([^)]*\)|_[A-Za-zΑ-ω]+|\^[A-Za-z0-9]+)")
    for part in pattern.split(text):
        if not part:
            continue
        special = pattern.fullmatch(part)
        value = part[1:] if special else part
        if special and value.startswith("{"):
            value = value[1:-1]
        run = p.add_run(value)
        font(run)
        if special:
            run.font.subscript = part[0] == "_"
            run.font.superscript = part[0] == "^"


def add_markdown(document, text):
    from docx.shared import Cm
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("| ") or line.startswith("|---"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i]); i += 1
            add_table(document, rows)
            continue
        if line.startswith("# "):
            paragraph(document, line[2:], heading=1)
        elif line.startswith("## "):
            paragraph(document, line[3:], heading=2)
        elif line.startswith("<equation>"):
            add_equation(document, line.replace("<equation>", "").replace("</equation>", ""))
        elif line.startswith("!["):
            match = re.fullmatch(r"!\[(.*?)\]\((.*?)\)", line)
            if not match:
                raise ValueError("无效图形标记")
            path = workspace(ROOT / match[2])
            if path.suffix.lower() == ".svg":
                path = render_svg(path)
            p = document.add_paragraph()
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(path), width=Cm(15.9))
            paragraph(document, match[1], centered=True)
        else:
            paragraph(document, line)
        i += 1


def create_paper(source, abstract, path, toc_pages=None):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
    document = setup_document()
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    font(p.add_run(TITLE), "黑体", 16, True)
    paragraph(document, "摘 要", heading=1)
    for text in abstract.split("\n\n"):
        paragraph(document, text)
    paragraph(document, "关键词：交通数字孪生；GIS；路径选择；源端协调；时间外推验证")
    document.add_page_break()
    paragraph(document, "目 录", heading=1)
    headings = re.findall(r"^# (.+)$", source, re.M)
    for heading in headings:
        p = document.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(0)
        font(p.add_run(heading + "  ……  " + str((toc_pages or {}).get(heading, "待更新"))))
    document.add_page_break()
    add_markdown(document, source)
    document.save(path)


def create_ai_document(source, path):
    document = setup_document()
    document.core_properties.title = "AI工具使用详情说明"
    add_markdown(document, source)
    document.save(path)


def convert_pdf(docx_path):
    font_config = OUT / "fontconfig.xml"
    font_config.write_text('''<?xml version="1.0"?><!DOCTYPE fontconfig SYSTEM "fonts.dtd"><fontconfig><include>/etc/fonts/fonts.conf</include><alias><family>宋体</family><prefer><family>Noto Serif CJK SC</family></prefer></alias><alias><family>SimSun</family><prefer><family>Noto Serif CJK SC</family></prefer></alias><alias><family>黑体</family><prefer><family>Noto Sans CJK SC</family></prefer></alias><alias><family>SimHei</family><prefer><family>Noto Sans CJK SC</family></prefer></alias><cachedir>''' + str(OUT / ".font-cache") + "</cachedir></fontconfig>", encoding="utf-8")
    with tempfile.TemporaryDirectory(prefix=".lo-profile-", dir=OUT) as profile, tempfile.TemporaryDirectory(prefix=".lo-temp-", dir=OUT) as temporary:
        env = os.environ.copy()
        env.update(FONTCONFIG_FILE=str(font_config), TMPDIR=temporary, XDG_CACHE_HOME=temporary, GSETTINGS_BACKEND="memory")
        result = subprocess.run(["/usr/bin/libreoffice", "-env:UserInstallation=" + Path(profile).as_uri(), "--headless", "--convert-to", "pdf:writer_pdf_Export", "--outdir", str(OUT), str(docx_path)],
                                env=env, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, timeout=180)
        target = OUT / (docx_path.stem + ".pdf")
        if result.returncode or not target.exists():
            raise RuntimeError(f"LibreOffice转换失败（退出码{result.returncode}）：" + result.stdout[-500:] + result.stderr[-500:])
        return target


def pdf_pages(path):
    result = subprocess.run(["/usr/bin/pdftotext", "-layout", str(path), "-"], capture_output=True, text=True, check=True)
    return result.stdout.rstrip("\f\n").split("\f")


def locate_headings(pages, headings):
    found = {}
    for index, page in enumerate(pages, 1):
        if index <= 2:
            continue
        normalized = [re.sub(r"\s+", "", line) for line in page.splitlines()]
        for heading in headings:
            if re.sub(r"\s+", "", heading) in normalized and heading not in found:
                found[heading] = index
    if set(found) != set(headings):
        raise ValueError("无法确定目录页码：" + str(set(headings) - set(found)))
    return found


def build_ai_text():
    return """# AI工具使用详情说明

本说明随支撑材料单独提供，不含队员姓名、学校、账户信息、密钥或签名。它记录当前可核查的AI参与事实，不代替参赛队员应独立承担的核心建模、仿真实现与分析责任。

## 1 工具、版本和使用日期

工具名称：OpenAI Codex；开发机构：OpenAI；使用日期：2026年10月17日至18日。当前会话未提供可核验的具体模型快照、精确版本号和版本发布日期，故未虚构填写。提交前应由队员从实际使用界面或账户记录补充可核验信息；无法取得时继续如实标为未知，并向主办方核实适用要求。

## 2 应用环节与实际贡献

| 环节 | AI实际参与 | 可核查产物 |
|---|---|---|
| 赛题与规范 | 阅读本地题目、评分、格式和AI规定；提取提交边界 | 论文方法映射、提交说明 |
| 数据处理 | 官方接口调用程序、脱敏、字段标准化、时间桶与排除原因 | 采集清单、处理脚本、冻结样本 |
| 拓扑与模型 | 方向边/转向审计、候选筛选、SUMO输入与运行诊断 | 路线目录、模型源码及留档 |
| 治理与实验 | A选路、B有界源端协调、对照实验及指标计算 | 最终协议、控制日志、汇总及单元测试 |
| 验证与展示 | 冻结观测层模型、留出误差、图表和离线回放 | 验证摘要、CSV、SVG/PNG及HTML |
| 论文与材料 | 章节草稿、引用核查、DOCX/PDF生成及格式检查 | 稿源、生成脚本、构建清单 |

AI参与包含核心技术程序与分析草稿，不仅是语句润色。原始百度响应由实际接口取得；仿真结果来自实际SUMO运行，论文数值从最终汇总文件读取；没有以AI直接生成的数值替代未执行实验。

## 3 关键交互摘要

以下为脱敏后的关键任务与响应摘要，非逐字完整聊天转录；不含任何鉴权参数。

| 交互主题 | 用户/任务约束摘要 | AI响应与实际采纳 |
|---|---|---|
| 工作区与赛题 | 读取工作区赛题；只在工作区修改，非必要不改变环境 | 只读赛题/规则；后续新增环境仅在工作区 |
| 正式数据源 | 研究上海道路，使用百度官方源 | 采集瑞金医院周边路线/路况，历史北京测试排除 |
| 采集时间 | 17日要求到16:15停止；18日要求再采集两小时 | 分会话保存实际截止、状态与来源清单 |
| 路径可行性 | 不虚构逆向边或转向；备选不增加需求 | 区分6个需求OD、拓扑证据与8条完整候选 |
| 物理检查 | 碰撞、传送或丢弃不得当成合格治理结果 | 保留失败记录，修复让行与内部连接后独立重跑 |
| 完成项目 | 完成基准、治理A/B、验证、论文和可复现材料 | 实际运行冻结数据验证和多情景对照，如实报告负结果 |

## 4 采纳内容、修改与核验

已采纳AI辅助形成的程序框架、质量检查、指标分解、图表和文稿。可查记录表明用户明确决定了城市、官方来源、工作区限制及采集时间；尚无证据证明参赛队员已经逐行审核核心代码、独立重建模型或完成对全部研究结论的人工修改。因此人工复核状态记录为“尚待队员完成”，不捏造人员签字或修改历史。

已执行的自动化检查包括原始输入摘要、路径连续性、候选简单性、同条件需求、源端等待计费、因果时序、输出守恒、碰撞/传送/丢弃、汇总对账以及留出窗口防泄漏。代理之间的交叉审阅属于AI辅助流程，不被称为独立的人类同行评审。检查通过仅证明其覆盖范围内的程序一致性，不能证明现实交通效果或替代参赛队的责任。

## 5 已知边界与提交前人工事项

真实数据仅为官方平台路线及路况快照，ETA不是独立逐车实测；仿真需求、道路车道参数、慢行事件和源端控制权限包含公开假设。拟合发生在观测层，不能改写成实际道路容量或全部物理参数已标定。若实验中B没有任务收益或增加普通车成本，论文继续保留该结果，不因期望正收益而删除。

队员需完成：独立理解模型与实现；核验方法和全部结果；必要实质修订并登记修改内容；核实工具版本与适用AI规定；检查论文匿名和本机字体；填写队号；按原模板独立签署承诺书。当前文件不表示这些人工事项已完成。赛题包AI规定题名为2025年试行，而论文格式规范提到2026年试行；应向主办方核对适用版本。

## 6 安全和复现记录

论文、AI详情和支撑包不包含真实姓名、学校、账户密钥、签名或含凭据URL。历史对话可能含私密信息，不直接导出完整聊天；本说明仅提供上述必要脱敏摘要。代码的AI辅助注释保留工具名称和已知日期；未提供的模型快照及发布日期标为未知。全部必要来源和构建摘要随成果保存，供队员审查。
"""


def submission_notes():
    return """# 提交前核对说明

稳定成果：outputs/paper/paper.pdf、paper.docx、ai_usage.pdf、ai_usage.docx。论文正文匿名，不含承诺书。DOCX用于队员实质修订与字体核验；PDF是当前可阅读版本。

论文与支撑材料截止北京时间2026年10月20日20:00；视频截止10月21日20:00。论文PDF和支撑ZIP/RAR分别不超过20MB；视频MP4、3分钟内、至少1280×720、不超过300MB。原文件未发现论文总页数上限。

正式命名：B队伍号.pdf；B队伍号fj.zip；承诺书B队伍号cns.pdf/jpg/png，另行由队员签名。队号未提供，不能猜填。视频格式规范与提交流程的连字符写法不一致，应按系统要求核对。上传后必须点击提交并确认“已提交”；本项目未自动上传或代签。

字体：DOCX指定题目三号黑体、一级四号黑体、其余中文小四宋体、单倍行距，无页眉，摘要起连续居中页码。当前系统没有正版SimSun/SimHei，PDF使用已有Noto宋体/黑体风格字体嵌入，构建清单如实记录。提交前必须在具有所需字体的合法环境打开DOCX检查并重新导出PDF；不能把当前字体替代称为已满足严格字体名称要求。

AI：工具已深度参与技术程序、分析和写作；队员独立理解、核心工作复核与实质修订尚未有完成人工记录，必须如实处理并按规定声明。AI具体模型快照/发布日期未知，不得编造；2025/2026版本差异需核对。AI详情PDF置入支撑包。

支撑包保留project.py、scripts/、src/、公开config/、必要脱敏data/raw和processed、冻结结果与关键图；完整大型XML可由脚本重生。不得纳入私密*.local.json、环境、缓存、competition规则或承诺书。论文附录列支撑结构；实际包内清单还须与之对应。
"""


def dynamic_content(validation, experiment, data):
    if experiment.get("protocol_sha256") != "ce9492e9f1164d11983fc5bf1c0e5e293fdbbc3f8053f6c7cb3a44b8360c0d1e" or data.get("snapshot_id") != "DATA_c6680fa324491ca5":
        raise ValueError("当前论文解释只适用于已冻结协议和数据；新实验须先实质修订稿件")
    if validation.get("status") != "eta_layer_evaluated":
        raise ValueError("观测层验证尚未完成")
    if validation["data_snapshot_id"] != data["snapshot_id"]:
        raise ValueError("验证与冻结数据不对应")
    if validation["input_hashes"]["data_manifest_sha256"] != digest(ROOT / "data/processed/final/data_manifest.json"):
        raise ValueError("数据摘要已改变")
    for artifact in validation["artifacts"].values():
        if digest(workspace(ROOT / artifact["path"])) != artifact["sha256"]:
            raise ValueError("验证产物摘要不一致")
    manifest = load(ROOT / "outputs/final_experiments/manifest.json")
    for name, expected in manifest["artifacts"].items():
        if digest(workspace(ROOT / "outputs/final_experiments" / name)) != expected:
            raise ValueError("实验产物摘要不一致：" + name)
    runs = experiment["runs"]
    if len(runs) != 22 or any(not r["physics_valid"] or r["unfinished_total"] for r in runs):
        raise ValueError("最终实验数量或完整性与论文协议不同")
    for r in runs:
        if digest(ROOT / r["folder"] / "run_manifest.json") != r["run_manifest_sha256"]:
            raise ValueError("运行摘要不一致：" + r["case"])
        if abs(r["ordinary_system_time_s"] - sum(r[k] for k in ("ordinary_in_network_time_s", "ordinary_source_control_wait_s", "ordinary_sumo_insertion_wait_s"))) > 1e-6:
            raise ValueError("完整成本无法对账")
        run_manifest=load(ROOT / r["folder"] / "run_manifest.json")
        for filename in ("trajectory.json", "source_queue_trace.json", "metrics.json"):
            if digest(ROOT / r["folder"] / filename) != run_manifest["artifacts"][filename]:
                raise ValueError("使用中的时序或指标文件摘要不一致")
    catalog = load(ROOT / "data/processed/baidu/route_catalog_v1.json")
    network_figure = make_network_figure(catalog)
    series_figure = make_validation_figure(validation)
    temporal_figure = make_experiment_figure()
    phases = [("training", "训练", "10月17日14:55—15:55"), ("same_day_exploratory", "同日探索", "10月17日15:55—16:15"),
              ("cross_day_holdout", "跨日留出", "10月18日11:36—13:36")]
    audit = data["data_audit"]
    traffic_count=sum(c["counts"]["traffic"] for c in data["collections"])
    dtext = (f"正式持续采集共取得{audit['raw_direction_count']}条路线响应，按统一规则移除{audit['duplicate_count']}条重复后，保留{audit['selected_direction_count']}条固定几何样本。持续观测覆盖北至南、南至北、西至医院外围和东至医院外围四个OD；六OD仿真需求与四OD连续观测不可混为同一数量。正式会话另有{traffic_count}次矩形路况响应，用于状态浏览及来源核查，未拿它们反演车流量。\n\n" +
        table(["分区", "北京时间（2026年）", "每OD条数", "合计"], [[label, window, str(audit["phases"][key]["eligible_count"] // 4), str(audit["phases"][key]["eligible_count"])] for key,label,window in phases]) +
        "\n\n全部160条样本的源观测时刻字段均缺失；有效样本中未发现路径变化、缺ETA或越界，接口失败数为0。请求记录对应17日14:57起至16:13止以及18日11:36起至13:34止的多种接口访问；路线最后一轮早于各会话最后路况请求。分区边界定义评价窗口，不等于每个边界瞬间都有采样。\n\n原计划17日另留30分钟未完整采齐，故只把已有20分钟标为探索；18日两小时单独评估。最终分区协议是在两个日期数据已经采集后整理冻结，不是事前注册实验。拟合和模型选择均不读取留出结果，但这一设计仍存在回顾性研究局限。")
    methods = [("physical_sumo", "未标定SUMO"), ("per_od_training_median", "OD训练中位数"), ("shared_affine_eta", "共享仿射")]
    metric_rows = []
    for key,label,_ in phases:
        for method,method_label in methods:
            m = validation["metrics"][key][method]
            metric_rows.append([label, method_label, n(m["sample_count"],0), n(m["mae_s"]), n(m["rmse_s"])])
    alpha = validation["models"]["shared_affine"]["alpha"]
    beta = validation["models"]["shared_affine"]["beta_s"]
    vtext = (f"本节拟合、误差统计和图表代码由AI辅助实现，数值来自实际冻结样本与SUMO输出。训练得到α={alpha:.6f}、β={beta:.3f}秒，非负约束未激活。四种不同T_k形成秩为2的设计矩阵，足以拟合两个统计参数，但不能识别更多物理机制。默认观测状态固定为OD训练中位数，共享仿射作为研究比较项；未根据留出误差重新选型或拟合。\n\n" +
        table(["OD", "SUMO在网均时/s", "训练中位数/s", "等效残差/s"], [[OD_NAMES[k],n(m["physical_sumo_s"]),n(m["training_median_s"]),n(m["equivalent_route_residual_s"])] for k,m in validation["models"]["per_od"].items()]) +
        "\n\n上表等效残差是平台中位数减固定假设下SUMO均时，混合了未识别信号、流量、几何和平台估计差异，不能解释为实测排队时间。\n\n" +
        table(["数据分区", "方法", "样本数", "MAE/s", "RMSE/s"], metric_rows) +
        "\n\n共享式在同日探索块的MAE为105.00秒，较中位数参照119.75秒低14.75秒；但跨日MAE为217.73秒，较中位数200.46秒高17.27秒，并且所有跨日预测均高估。未标定SUMO跨日MAE为200.60秒，其接近中位数误差并不说明内部物理状态正确：前者普遍低估，后者普遍高估。结果说明单时段静态映射没有稳定的跨日优势，不能宣称泛化验证成功。\n\n" +
        f"![图2 跨日平台ETA与冻结参照。AI辅助绘图；点为实际请求对应的约5分钟快照。]({series_figure.relative_to(ROOT)})\n\n" +
        "图2显示18日上午至午后的平台ETA整体低于17日下午训练状态。日期和时段同时变化，现有样本无法把周末日差异与日内变化分离，也不支持将误差下降归因于单一物理原因。")
    by_case = {r["case"]:r for r in runs}
    profile_order = ["primary", "higher_demand", "severe_event", "longer_event"]
    etext = "本节实验运行、统计和解释草稿由AI辅助完成，输入与实际SUMO产物均留档。最终独立运行22组：2组S1、8组S2和12组S3。各组全部需求到达，未完成、碰撞、传送、丢弃均为0；比较前核验需求、事件及物理参数摘要。下表单位“车秒”是所有普通车完整时间之和，任务时间另列，避免将不同目标合成任意权重的单一分数。\n\n"
    etext += table(["S1设置", "普通车辆数", "完整时间/车秒", "入网等待/车秒"], [[PROFILES[r["profile"]],n(r["ordinary_planned"],0),n(r["ordinary_system_time_s"],0),n(r["ordinary_sumo_insertion_wait_s"],0)] for r in runs if r["scenario_family"]=="S1"])
    etext += "\n\nS2固定路径与A对照（所有组均计入等待）：\n\n" + table(["设置", "固定/车秒", "A/车秒", "A−固定", "A改道数"], [[PROFILES[p],n(by_case[p+"_S2_fixed"]["ordinary_system_time_s"],0),n(by_case[p+"_S2_A"]["ordinary_system_time_s"],0),n(by_case[p+"_S2_A"]["ordinary_system_time_s"]-by_case[p+"_S2_fixed"]["ordinary_system_time_s"],0),n(by_case[p+"_S2_A"]["changed_route_count"],0)] for p in profile_order])
    etext += "\n\nS3全部策略结果（普通车数为主设置240辆、需求加倍480辆，另加同一任务1辆）：\n\n" + table(["设置", "策略", "普通时间/车秒", "任务时间/s", "源端等待/车秒"], [[PROFILES[r["profile"]],{"fixed":"固定","A":"A","B":"B"}[r["policy"]],n(r["ordinary_system_time_s"],0),n(r["task_release_to_arrival_s"],0),n(r["ordinary_source_control_wait_s"],0)] for r in runs if r["scenario_family"]=="S3"])
    etext += "\n\nB相对A的边际效应：\n\n" + table(["设置", "普通Δ/车秒", "任务Δ/s", "受控车数", "最大源端等待/s"], [[PROFILES[p],n(by_case[p+"_S3_B"]["ordinary_system_time_s"]-by_case[p+"_S3_A"]["ordinary_system_time_s"],0),n(by_case[p+"_S3_B"]["task_release_to_arrival_s"]-by_case[p+"_S3_A"]["task_release_to_arrival_s"],0),n(by_case[p+"_S3_B"]["source_held_vehicle_count"],0),n(by_case[p+"_S3_B"]["maximum_source_control_wait_s"],0)] for p in profile_order])
    fairness = load(ROOT / "outputs/quality/fairness_audit.json")
    etext += "\n\n逐车分布由同一车辆编号和原始释放时刻对账，45秒上限不保证任何车辆总损失≤45秒：\n\n" + table(["设置", "变快/辆", "变慢/辆", "不变/辆", "最大总损失/s"], [[PROFILES[r["profile"]],n(r["benefited_count"],0),n(r["harmed_count"],0),n(r["unchanged_count"],0),n(r["maximum_individual_loss_s"],0)] for r in fairness["profiles"]])
    etext += "\n\n高需求组普通交通总体减少3340车秒，仍有30辆变慢、127辆变快；主设置和长事件组单车最大总损失58秒。本文只称源端控制等待有界与总体条件性改善，不称帕累托改善。"
    interpretation = ("A只在0.5 m/s的更强慢行情景改变路线：S2普通总时间49071降至44766车秒，减少4305车秒（8.77%）；S3固定至A，普通总时间50226降至44774车秒，任务323降至200秒。1 m/s主设置、需求加倍与事件延长下，当前外部边成本仍选择默认路径，A均无收益。这符合长绕行与局部慢行成本之间的权衡，也暴露了不预测排队的局限。\n\n"
        "B在高需求S3相对A使任务292降至232秒，减少60秒（20.55%）；普通完整时间108374降至105034车秒，减少3340车秒（3.08%）。其代价分解为：源端等待0增至1887车秒，入网等待7873降至7398车秒，在网时间100501降至95749车秒。源端等待没有被删除，任务两组均在720秒实际入网，到达1012秒和952秒，收益并非提前发车。\n\n"
        "主设置B未改善174秒任务时间，却增加995车秒普通成本；更强慢行中B使任务200增至201秒、普通增加1000车秒；事件延长亦增加995车秒而任务不变。模型内一种合理解释是，高需求下延后部分入流降低了任务走廊的后续竞争；较低需求时相同控制偏强，额外等待与恢复放行团簇反而造成代价。这里是仿真机制解释，不是现场因果结论。\n\n" +
        f"![图3 AI辅助统计实际5秒仿真输出：上为在网车辆数，下为B源端等待数；横轴不表示实测交通时间。]({temporal_figure.relative_to(ROOT)})\n\n" +
        "图3保留事件开始600秒、任务发布720秒和控制窗口结束900秒的时序。在任务发布前，A/B的FCD状态逐项一致；发布后路线分配相同，只改变许可源端放行。该逐步核验支持将本轮A/B差异归于已实现控制机制，但不排除真实系统中尚未建模的其他影响。")
    conclusion = ("本文完成了官方GIS证据、方向拓扑、观测层检验和微观治理对照的统一原型。160条有效平台路线快照支持可追溯时间分区，307个有向外部分段、310个允许转向和8条候选支持六OD需求下的合法动作审计，22组独立运行均通过所声明物理检查。\n\n"
        "结果具有明确条件性：单时段共享ETA映射跨日误差大于中位数参照；A在更强慢行下有收益而其余设置保持原路；B仅在高需求组合使任务缩短60秒、普通总成本减少3340车秒，其他组的无益或不利结果保留。源端等待有界不等于个体总损失有界，也不意味着所有普通车受益。\n\n"
        "研究贡献是可执行、可对账且不隐藏负结果的方法链。下一步应取得独立逐车时间、实测流量、信号及入口证据，扩充时间窗口并校准物理模型，再讨论现场部署。当前数值只适用于公开假设与已观测子图，不作为现实上海道路治理效果承诺。")
    time_mapping = "为统一界面显示，最终实验声明t=0对应2026年10月18日12:00:00（北京时间），所以600、720、1200、1800秒可分别标作12:10、12:12、12:20、12:30。这是反事实实验时钟约定，不是从同一时段采集重建出的车流、事故或急救任务；展示必须保留该标记。真实快照仍按自身请求/接收时刻展示，不以实验锚点覆盖。仿真图中横轴优先使用秒，任务到达以实际SUMO输出为准。"
    refs = load(SOURCE / "references.json")
    blocks = {"DATA_RESULTS":dtext,"NETWORK_FIGURE":f"![图1 AI辅助绘制实际GIS子图、西侧任务路线与假设事件段；箭头依据源路线方向。]({network_figure.relative_to(ROOT)})",
              "VALIDATION_RESULTS":vtext,"TIME_MAPPING":time_mapping,"EXPERIMENT_RESULTS":etext,"RESULT_INTERPRETATION":interpretation,"CONCLUSION":conclusion,
              "REFERENCES":"\n\n".join(f"[{r['id']}] {r['text']}" for r in refs)}
    source = (SOURCE / "paper_core.md").read_text(encoding="utf-8")
    for key, value in blocks.items():
        source = source.replace("{{"+key+"}}", value)
    abstract = ("针对城市街区中道路状态可见而干预效果难以核验的问题，建立官方GIS证据、方向拓扑、路线观测与微观仿真相结合的交通数字孪生原型。以上海瑞金医院周边为例，实时归档百度路线与路况接口，按五分钟桶保留160条有效固定路线快照；建立含307个有向外部分段、310个允许转向、六个需求OD和八条完整候选的局部路网。将地图几何、平台预计时间与仿真轨迹分层保存，避免以路线估计代替逐车实测。\n\n"
        "在一般交通基准上，只用训练块拟合非负共享ETA映射，并保留OD中位数参照。跨日96条留出样本中，共享映射MAE为217.73秒，高于参照200.46秒，显示单时段状态不能直接跨日迁移。进一步提出当前信息选路A，以及任务发布后的有界源端协调B，完成22组固定协议仿真，全部到达且碰撞、传送、丢弃为零。更强慢行下A使普通交通总时间减少8.77%；高需求S3中，B相对A使任务292降至232秒、普通交通总成本减少3340车秒，但仍有30辆普通车变慢。主设置B无任务收益，普通成本增加995车秒。\n\n"
        "结果表明路径可行性、因果时序和完整等待计费共同决定治理结论。本文保留无收益与不利结果，给出可复现输入、输出和离线回放；平台ETA观测层检验不等同于现场物理校准，假设事件及源端控制效果限于当前子图和参数。")
    return source, abstract


def main(argv=None):
    parser = argparse.ArgumentParser(description="从真实最终摘要生成匿名论文和AI说明，使用工作区.venv内python-docx及现有LibreOffice")
    parser.add_argument("--validation", type=Path, default=ROOT / "outputs/validation/summary.json")
    parser.add_argument("--experiments", type=Path, default=ROOT / "outputs/final_experiments/summary.json")
    parser.add_argument("--data", type=Path, default=ROOT / "data/processed/final/data_manifest.json")
    args = parser.parse_args(argv)
    OUT.mkdir(parents=True, exist_ok=True)
    validation, experiment, data = map(load, (args.validation, args.experiments, args.data))
    if experiment.get("status") != "completed":
        raise ValueError("最终仿真实验尚未完成，拒绝填造论文结果")
    source, abstract = dynamic_content(validation, experiment, data)
    if re.search(r"\{\{[A-Z_]+\}\}", source):
        raise ValueError("稿件仍有未填入真实结果的标记")
    (SOURCE / "paper_source.md").write_text(source, encoding="utf-8")
    ai = build_ai_text()
    (SOURCE / "ai_usage.md").write_text(ai, encoding="utf-8")
    (SOURCE / "submission_notes.md").write_text(submission_notes(), encoding="utf-8")
    create_paper(source, abstract, OUT / "paper.docx")
    pdf = convert_pdf(OUT / "paper.docx")
    headings = re.findall(r"^# (.+)$", source, re.M)
    page_numbers = locate_headings(pdf_pages(pdf), headings)
    create_paper(source, abstract, OUT / "paper.docx", page_numbers)
    pdf = convert_pdf(OUT / "paper.docx")
    pages = pdf_pages(pdf)
    if locate_headings(pages, headings) != page_numbers:
        raise ValueError("更新目录后页码变化，须重新排版核验")
    create_ai_document(ai, OUT / "ai_usage.docx")
    ai_pdf = convert_pdf(OUT / "ai_usage.docx")
    for path in (pdf, ai_pdf):
        if path.stat().st_size > 20_000_000:
            raise ValueError("PDF超过20MB")
    font_output = subprocess.run(["pdffonts", str(pdf)], capture_output=True, text=True, check=True).stdout
    manifest = {"generated_at": datetime.now(timezone.utc).isoformat(), "paper_pages": len(pages),
        "heading_pages": page_numbers, "anonymous_metadata": True,
        "font_status": "DOCX指定宋体黑体；本机缺字体，PDF采用Noto宋体/黑体风格替代；提交前须合法字体环境核对",
        "actual_pdf_fonts": font_output, "human_review_status": "not_recorded_complete",
        "source_sha256": {str(workspace(p).relative_to(ROOT)): digest(workspace(p)) for p in (args.validation, args.experiments, args.data, Path(__file__), SOURCE / "paper_core.md", SOURCE / "references.json", ROOT / "outputs/quality/fairness_audit.json")},
        "artifacts": {str(p.relative_to(OUT)): {"sha256": digest(p), "bytes": p.stat().st_size} for p in (OUT / "paper.docx", pdf, OUT / "ai_usage.docx", ai_pdf, *sorted((OUT / "figures").glob("*.png")))}}
    (OUT / "build_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"paper": str(pdf.relative_to(ROOT)), "pages": len(pages), "AI_details": str(ai_pdf.relative_to(ROOT)), "font_review_required": True}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
